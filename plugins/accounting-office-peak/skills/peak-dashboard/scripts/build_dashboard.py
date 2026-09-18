#!/usr/bin/env python3
"""สร้าง Dashboard .xlsx จาก data pack (JSON) ที่ดึงมาจาก ERP

เหตุผลที่เป็นสคริปต์ ไม่ใช่คำสั่งใน SKILL.md:
ตัวเลขในรายงานผู้บริหารต้องออกมาเหมือนกันทุกครั้งที่รันด้วยข้อมูลชุดเดิม
ถ้าให้โมเดลเขียนโค้ด openpyxl ใหม่ทุกครั้ง หน้าตาไฟล์และสูตรจะเปลี่ยนไปเรื่อย ๆ
และการ "ดึงไม่ได้" จะกลายเป็น 0 โดยไม่มีใครรู้ — ซึ่งเป็นความผิดพลาดที่แพงที่สุดของงานนี้

สคริปต์นี้ไม่ต่อเน็ตและไม่รู้จัก ERP — มันรู้แค่รูปของ data pack
หน้าที่ดึงข้อมูลเป็นของ skill ดูสัญญาของไฟล์ JSON ใน references/dashboard-layout.md

    python3 build_dashboard.py <data.json> --out <ไฟล์.xlsx> [--plugin-root <dir>]

ต้องมี openpyxl ถ้าไม่มีจะบอกวิธีติดตั้งแล้วออกด้วย exit code 3 (ไม่ใช่ traceback)

สำเนาไฟล์นี้ต้องตรงกันทั้ง 2 plugin ของ ERP — ดู scripts/sync_hooks_lib.py
"""
from __future__ import annotations

import argparse
import datetime
import json
import os
import sys

# ---------------------------------------------------------------- สี (ใช้ชุดเดียวกับ skill อื่นในชุดนี้)
GREEN = "C6EFCE"      # ดีขึ้น / ผ่าน
RED = "FFC7CE"        # แย่ลง / ไม่ผ่าน
YELLOW = "FFEB9C"     # ต้องดู
GREY = "D9D9D9"       # ดึงไม่ได้ — ไม่ใช่ศูนย์
HEAD_BG = "1F4E79"
HEAD_FG = "FFFFFF"
TITLE_BG = "DDEBF7"

MONEY_FMT = "#,##0.00"
INT_FMT = "#,##0"
PCT_FMT = "0.0%"

UNAVAILABLE = "ดึงไม่ได้"


def die(msg, code=2):
    sys.stderr.write("build_dashboard: %s\n" % msg)
    raise SystemExit(code)


try:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:                                     # pragma: no cover
    die("ต้องมี openpyxl — ติดตั้งด้วย `python3 -m pip install --user openpyxl` แล้วรันใหม่", 3)


# ---------------------------------------------------------------- หัวไฟล์ (กฎบังคับข้อ 5)
def plugin_meta(root=None):
    """อ่านชื่อ+เวอร์ชันจาก .claude-plugin/plugin.json — ห้าม hardcode เลขเวอร์ชัน"""
    root = root or os.environ.get("CLAUDE_PLUGIN_ROOT")
    if not root:
        # scripts/ -> <skill>/ -> skills/ -> <plugin root>
        root = os.path.dirname(os.path.dirname(os.path.dirname(
            os.path.dirname(os.path.abspath(__file__)))))
    try:
        with open(os.path.join(root, ".claude-plugin", "plugin.json"), encoding="utf-8") as fh:
            data = json.load(fh)
        return data.get("name") or "accounting-office", data.get("version") or "0.0.0"
    except Exception:
        return "accounting-office", "0.0.0"


def thai_today():
    d = datetime.date.today()
    return "%04d-%02d-%02d" % (d.year + 543, d.month, d.day)


# ---------------------------------------------------------------- ตัวช่วยเล็ก ๆ
def money(value):
    return value if isinstance(value, (int, float)) else None


def fill(color):
    return PatternFill("solid", fgColor=color)


THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def put_header_row(ws, row, labels, widths=None):
    for i, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=label)
        c.font = Font(bold=True, color=HEAD_FG)
        c.fill = fill(HEAD_BG)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def put_title(ws, row, text, span=6):
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(bold=True, size=12)
    c.fill = fill(TITLE_BG)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max(span, 2))
    return row + 1


def status_of(block):
    st = str((block or {}).get("status") or "ok").lower()
    return st if st in ("ok", "unavailable", "partial") else "ok"


def unavailable_note(block):
    reason = (block or {}).get("reason") or "ไม่ได้ระบุเหตุผล"
    return "%s — %s" % (UNAVAILABLE, reason)


def source_text(block):
    src = (block or {}).get("source") or {}
    tool = src.get("tool") or "(ไม่ระบุ tool)"
    params = src.get("params")
    if isinstance(params, dict) and params:
        params = " ".join("%s=%s" % (k, v) for k, v in params.items())
    return "%s | %s" % (tool, params or "-")


# ---------------------------------------------------------------- เขียนบล็อกเป็นตาราง
def write_table(ws, block, start_row):
    """คืน dict บอกตำแหน่งข้อมูลที่เขียนไว้ เพื่อให้กราฟอ้างถึงได้"""
    title = block.get("title") or block.get("id") or ""
    row = put_title(ws, start_row, title, span=max(2, len(block.get("columns") or [])))

    note = ws.cell(row=row, column=1, value="ที่มา: %s" % source_text(block))
    note.font = Font(italic=True, size=9, color="595959")
    row += 1

    columns = block.get("columns") or []
    if not columns:
        die("บล็อก %r ไม่มี columns" % block.get("id"))

    st = status_of(block)
    if st == "unavailable":
        c = ws.cell(row=row, column=1, value=unavailable_note(block))
        c.fill = fill(GREY)
        c.font = Font(bold=True)
        ws.merge_cells(start_row=row, start_column=1, end_row=row,
                       end_column=max(2, len(columns)))
        return {"id": block.get("id"), "sheet": ws.title, "status": st,
                "rows": 0, "next_row": row + 2}

    labels = [col.get("label") or col.get("key") for col in columns]
    widths = [col.get("width") or (14 if col.get("type") in ("money", "int", "pct") else 34)
              for col in columns]
    put_header_row(ws, row, labels, widths)
    header_row = row
    row += 1

    first_data = row
    rows = block.get("rows") or []
    for item in rows:
        for i, col in enumerate(columns, start=1):
            key = col.get("key")
            raw = item.get(key) if isinstance(item, dict) else None
            ctype = col.get("type") or "text"
            cell = ws.cell(row=row, column=i)
            if raw is None or raw == "":
                cell.value = "" if ctype == "text" else UNAVAILABLE
                if ctype != "text":
                    cell.fill = fill(GREY)
            elif ctype == "money":
                cell.value = money(raw)
                cell.number_format = MONEY_FMT
            elif ctype == "int":
                cell.value = raw
                cell.number_format = INT_FMT
            elif ctype == "pct":
                cell.value = raw
                cell.number_format = PCT_FMT
            else:
                cell.value = raw
            cell.border = BOX
            if ctype == "pct" and isinstance(raw, (int, float)):
                cell.fill = fill(GREEN if raw >= 0 else RED)
        row += 1

    last_data = row - 1
    if rows and block.get("total_row"):
        for i, col in enumerate(columns, start=1):
            cell = ws.cell(row=row, column=i)
            cell.font = Font(bold=True)
            cell.fill = fill(TITLE_BG)
            cell.border = BOX
            if i == 1:
                cell.value = "รวม"
            elif col.get("type") in ("money", "int") and col.get("total", True):
                letter = get_column_letter(i)
                cell.value = "=SUM(%s%d:%s%d)" % (letter, first_data, letter, last_data)
                cell.number_format = MONEY_FMT if col.get("type") == "money" else INT_FMT
        row += 1

    if st == "partial":
        c = ws.cell(row=row, column=1, value="⚠️ ข้อมูลไม่ครบ — %s"
                    % (block.get("reason") or "ดูชีต ที่มาของข้อมูล"))
        c.fill = fill(YELLOW)
        row += 1

    return {"id": block.get("id"), "sheet": ws.title, "status": st, "rows": len(rows),
            "header_row": header_row, "first_data": first_data, "last_data": last_data,
            "col_index": {col.get("key"): i for i, col in enumerate(columns, start=1)},
            "next_row": row + 1}


# ---------------------------------------------------------------- KPI
def write_kpis(ws, kpis, start_row):
    row = put_title(ws, start_row, "ตัวเลขสำคัญ", span=6)
    put_header_row(ws, row,
                   ["ตัวชี้วัด", "งวดนี้", "งวดเทียบ", "เปลี่ยนแปลง", "สถานะ", "ที่มา"],
                   [34, 16, 16, 14, 14, 46])
    row += 1
    for kpi in kpis or []:
        name = kpi.get("name") or "(ไม่ระบุชื่อ)"
        ws.cell(row=row, column=1, value=name).border = BOX
        st = status_of(kpi)
        if st == "unavailable":
            c = ws.cell(row=row, column=2, value=unavailable_note(kpi))
            c.fill = fill(GREY)
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
            ws.cell(row=row, column=5, value=UNAVAILABLE).fill = fill(GREY)
        else:
            v = money(kpi.get("value"))
            cv = money(kpi.get("compare"))
            c = ws.cell(row=row, column=2, value=v)
            c.number_format = MONEY_FMT
            c2 = ws.cell(row=row, column=3, value=cv)
            c2.number_format = MONEY_FMT
            c3 = ws.cell(row=row, column=4)
            if v is not None and cv:
                change = (v - cv) / abs(cv)
                c3.value = change
                c3.number_format = PCT_FMT
                good = kpi.get("higher_is_better", True)
                better = change >= 0 if good else change <= 0
                c3.fill = fill(GREEN if better else RED)
            else:
                c3.value = "-"
            ws.cell(row=row, column=5, value="ok" if st == "ok" else "ข้อมูลไม่ครบ").fill = \
                fill(GREEN if st == "ok" else YELLOW)
        ws.cell(row=row, column=6, value=source_text(kpi)).font = Font(size=9, color="595959")
        for col in range(1, 7):
            ws.cell(row=row, column=col).border = BOX
        row += 1
    return row + 1


# ---------------------------------------------------------------- การตรวจยอด
def write_checks(ws, checks, start_row):
    """ตรวจยอดข้ามแหล่ง — ไม่ปรับตัวเลขให้ตรงกัน แค่บอกว่าตรงหรือไม่ตรงเท่าไร"""
    row = put_title(ws, start_row, "การตรวจยอดข้ามแหล่ง", span=6)
    put_header_row(ws, row, ["รายการตรวจ", "ค่าจากแหล่ง ก", "ค่าจากแหล่ง ข",
                             "ผลต่าง", "เกณฑ์", "ผล"], [46, 16, 16, 14, 12, 14])
    row += 1
    failed = []
    for chk in checks or []:
        name = chk.get("name") or "(ไม่ระบุ)"
        left, right = money(chk.get("left")), money(chk.get("right"))
        tol = chk.get("tolerance")
        tol = tol if isinstance(tol, (int, float)) else 1.0
        ws.cell(row=row, column=1, value=name)
        for col, val in ((2, left), (3, right)):
            c = ws.cell(row=row, column=col)
            if val is None:
                c.value = UNAVAILABLE
                c.fill = fill(GREY)
            else:
                c.value = val
                c.number_format = MONEY_FMT
        cd = ws.cell(row=row, column=4)
        cr = ws.cell(row=row, column=6)
        ws.cell(row=row, column=5, value=tol).number_format = MONEY_FMT
        if left is None or right is None:
            cd.value = UNAVAILABLE
            cd.fill = fill(GREY)
            cr.value = "ตรวจไม่ได้"
            cr.fill = fill(GREY)
            failed.append("%s (ตรวจไม่ได้)" % name)
        else:
            diff = left - right
            cd.value = diff
            cd.number_format = MONEY_FMT
            ok = abs(diff) <= tol
            cr.value = "ผ่าน" if ok else "ไม่ผ่าน"
            cr.fill = fill(GREEN if ok else RED)
            if not ok:
                failed.append("%s (ต่าง %.2f)" % (name, diff))
        for col in range(1, 7):
            ws.cell(row=row, column=col).border = BOX
        row += 1
    return row + 1, failed


# ---------------------------------------------------------------- กราฟ
CHART_TYPES = {"bar": BarChart, "line": LineChart, "pie": PieChart}


def add_charts(wb, dash, specs, placed, first_row):
    """วางกราฟบนชีตแดชบอร์ด โดยอ้างเซลล์ในชีตของบล็อกนั้น (cross-sheet ref)"""
    skipped = []
    for spec in specs or []:
        block_id = spec.get("block")
        info = placed.get(block_id)
        if not info or info.get("status") == "unavailable" or not info.get("rows"):
            skipped.append("%s (ไม่มีข้อมูล)" % (spec.get("title") or block_id))
            continue
        ctype = CHART_TYPES.get(str(spec.get("type") or "bar").lower())
        if ctype is None:
            skipped.append("%s (ชนิดกราฟ %r ไม่รองรับ)" % (block_id, spec.get("type")))
            continue
        ws = wb[info["sheet"]]
        cat_col = info["col_index"].get(spec.get("cat"))
        val_keys = [k for k in (spec.get("values") or []) if k in info["col_index"]]
        if not cat_col or not val_keys:
            skipped.append("%s (คอลัมน์ที่อ้างไม่มีในบล็อก)" % block_id)
            continue
        chart = ctype()
        chart.title = spec.get("title") or info["id"]
        chart.height = float(spec.get("height") or 7.5)
        chart.width = float(spec.get("width") or 15)
        cats = Reference(ws, min_col=cat_col, min_row=info["first_data"],
                         max_row=info["last_data"])
        for key in val_keys:
            col = info["col_index"][key]
            data = Reference(ws, min_col=col, min_row=info["header_row"],
                             max_row=info["last_data"])
            chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        dash.add_chart(chart, spec.get("anchor") or ("H%d" % first_row))
    return skipped


# ---------------------------------------------------------------- ชีตที่มาของข้อมูล
def write_sources(ws, pack, placed):
    put_header_row(ws, 1, ["บล็อก", "ชีต", "tool ที่เรียก", "พารามิเตอร์", "จำนวนแถว",
                           "สถานะ", "ดึงเมื่อ", "หมายเหตุ"],
                   [24, 22, 42, 34, 12, 14, 20, 46])
    row = 2
    for block in pack.get("blocks") or []:
        src = block.get("source") or {}
        info = placed.get(block.get("id")) or {}
        params = src.get("params")
        if isinstance(params, dict):
            params = " ".join("%s=%s" % (k, v) for k, v in params.items())
        st = status_of(block)
        values = [block.get("id"), info.get("sheet") or block.get("sheet"),
                  src.get("tool") or "-", params or "-",
                  info.get("rows", len(block.get("rows") or [])),
                  {"ok": "ok", "partial": "ข้อมูลไม่ครบ", "unavailable": UNAVAILABLE}[st],
                  src.get("fetched_at") or "-", block.get("reason") or ""]
        for i, v in enumerate(values, start=1):
            c = ws.cell(row=row, column=i, value=v)
            c.border = BOX
            if i == 6:
                c.fill = fill({"ok": GREEN, "partial": YELLOW, "unavailable": GREY}[st])
        row += 1
    for kpi in pack.get("kpis") or []:
        src = kpi.get("source") or {}
        st = status_of(kpi)
        values = ["kpi:%s" % (kpi.get("name") or "?"), "แดชบอร์ด", src.get("tool") or "-",
                  src.get("params") if isinstance(src.get("params"), str) else
                  " ".join("%s=%s" % (k, v) for k, v in (src.get("params") or {}).items()) or "-",
                  "", {"ok": "ok", "partial": "ข้อมูลไม่ครบ", "unavailable": UNAVAILABLE}[st],
                  src.get("fetched_at") or "-", kpi.get("reason") or ""]
        for i, v in enumerate(values, start=1):
            c = ws.cell(row=row, column=i, value=v)
            c.border = BOX
            if i == 6:
                c.fill = fill({"ok": GREEN, "partial": YELLOW, "unavailable": GREY}[st])
        row += 1
    return row


# ---------------------------------------------------------------- main
def build(pack, out_path, plugin_root=None):
    name, version = plugin_meta(plugin_root)
    meta = pack.get("meta") or {}

    wb = Workbook()
    dash = wb.active
    dash.title = "แดชบอร์ด"

    header = "%s v%s | ลูกค้า: %s | กิจการใน ERP: %s | งวด: %s | เทียบกับ: %s | ดึงเมื่อ: %s" % (
        name, version,
        meta.get("client") or "-", meta.get("company_in_erp") or "-",
        meta.get("period_label") or "-", meta.get("compare_label") or "-",
        meta.get("fetched_at") or thai_today())
    c = dash.cell(row=1, column=1, value=header)
    c.font = Font(bold=True, size=11)
    dash.cell(row=2, column=1, value=(
        "ตัวเลขทั้งหมดดึงจาก %s ตามที่ระบบรายงาน ณ เวลาที่ดึง — ช่อง '%s' คือดึงไม่ได้ ไม่ใช่ศูนย์ "
        "· ดูที่มาของทุกตัวเลขในชีต 'ที่มาของข้อมูล'"
        % (meta.get("erp") or "ERP", UNAVAILABLE))).font = Font(italic=True, size=9,
                                                                color="595959")
    dash.column_dimensions["A"].width = 40
    for col in "BCDE":
        dash.column_dimensions[col].width = 16
    dash.column_dimensions["F"].width = 46

    row = write_kpis(dash, pack.get("kpis"), 4)
    chart_anchor_row = row
    checks_row = row + 22                      # เว้นที่ให้กราฟก่อนตารางตรวจยอด
    checks_row, failed_checks = write_checks(dash, pack.get("checks"), checks_row)

    if meta.get("notes"):
        row = put_title(dash, checks_row, "ข้อสังเกตจากผู้จัดทำ", span=6)
        for note in meta["notes"]:
            dash.cell(row=row, column=1, value="• %s" % note)
            row += 1

    placed = {}
    for block in pack.get("blocks") or []:
        sheet_name = (block.get("sheet") or block.get("id") or "ข้อมูล")[:31]
        ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.create_sheet(sheet_name)
        start = ws.max_row + 2 if ws.max_row > 1 else 1
        placed[block.get("id")] = write_table(ws, block, start)

    skipped = add_charts(wb, dash, pack.get("charts"), placed, chart_anchor_row)

    write_sources(wb.create_sheet("ที่มาของข้อมูล"), pack, placed)

    out_dir = os.path.dirname(os.path.abspath(out_path))
    if out_dir and not os.path.isdir(out_dir):
        os.makedirs(out_dir, exist_ok=True)
    wb.save(out_path)
    return {"path": out_path, "header": header, "placed": placed,
            "failed_checks": failed_checks, "skipped_charts": skipped}


def main(argv=None):
    ap = argparse.ArgumentParser(description="สร้าง Dashboard .xlsx จาก data pack (JSON)")
    ap.add_argument("data", help="ไฟล์ JSON ตามสัญญาใน references/dashboard-layout.md")
    ap.add_argument("--out", required=True, help="path ของไฟล์ .xlsx ที่จะเขียน")
    ap.add_argument("--plugin-root", help="ที่อยู่ของ plugin (ปกติอ่านจาก CLAUDE_PLUGIN_ROOT)")
    args = ap.parse_args(argv)

    if not os.path.isfile(args.data):
        die("ไม่พบไฟล์ข้อมูล %s" % args.data)
    try:
        with open(args.data, encoding="utf-8") as fh:
            pack = json.load(fh)
    except json.JSONDecodeError as exc:
        die("ไฟล์ข้อมูลไม่ใช่ JSON ที่อ่านได้: %s" % exc)

    if not isinstance(pack, dict) or not (pack.get("blocks") or pack.get("kpis")):
        die("data pack ต้องมีอย่างน้อย kpis หรือ blocks")

    result = build(pack, args.out, args.plugin_root)

    print("เขียนไฟล์แล้ว: %s" % result["path"])
    print(result["header"])
    unavailable = [i for i in result["placed"].values() if i.get("status") == "unavailable"]
    partial = [i for i in result["placed"].values() if i.get("status") == "partial"]
    print("บล็อกทั้งหมด %d · ดึงไม่ได้ %d · ข้อมูลไม่ครบ %d"
          % (len(result["placed"]), len(unavailable), len(partial)))
    for info in unavailable:
        print("  - %s: %s (ต้องบอกผู้ใช้ว่าช่องนี้ว่างเพราะดึงไม่ได้)" % (info["id"], UNAVAILABLE))
    if result["skipped_charts"]:
        print("กราฟที่ไม่ได้วาง: %s" % "; ".join(result["skipped_charts"]))
    if result["failed_checks"]:
        print("⚠️ การตรวจยอดไม่ผ่าน: %s" % "; ".join(result["failed_checks"]))
        print("   ห้ามปรับตัวเลขให้ตรงกันเอง — รายงานผลต่างตามจริงให้ผู้ใช้")
    return 0


if __name__ == "__main__":
    sys.exit(main())
