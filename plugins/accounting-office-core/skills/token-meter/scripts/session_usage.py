#!/usr/bin/env python3
"""Summarise Claude Code token usage and cost from session transcripts (JSONL).

Examples
  python session_usage.py                         # latest session in ~/.claude/projects
  python session_usage.py --all --since 2026-09-01
  python session_usage.py --session <id-or-file> --turns
  python session_usage.py --by-skill --xlsx out.xlsx --fx 32.5

A "turn" = one real user prompt plus every assistant call that follows it until
the next real prompt. A turn is attributed to every skill invoked inside it
(Skill tool call or /slash-command).
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from tm_common import cost_usd, fmt_thb, fmt_usd, load_pricing, price_for_model  # noqa: E402

KEYS = ("input", "output", "cache_write_5m", "cache_write_1h", "cache_read")
CMD_RE = re.compile(r"<command-name>\s*/?([^<\s]+)\s*</command-name>")


def projects_dir() -> Path:
    base = os.environ.get("CLAUDE_CONFIG_DIR") or str(Path.home() / ".claude")
    return Path(base) / "projects"


def norm_usage(u: dict) -> dict:
    cc = u.get("cache_creation") or {}
    w5 = cc.get("ephemeral_5m_input_tokens")
    w1 = cc.get("ephemeral_1h_input_tokens", 0) or 0
    if w5 is None:  # older transcripts only have the total
        w5 = (u.get("cache_creation_input_tokens") or 0) - w1
    return {
        "input": u.get("input_tokens") or 0,
        "output": u.get("output_tokens") or 0,
        "cache_write_5m": max(0, w5),
        "cache_write_1h": w1,
        "cache_read": u.get("cache_read_input_tokens") or 0,
    }


def is_real_prompt(msg: dict) -> bool:
    """User line typed by a human (not a tool_result echo)."""
    if msg.get("isMeta"):
        return False
    content = (msg.get("message") or {}).get("content")
    if isinstance(content, str):
        return True
    if isinstance(content, list):
        return any(b.get("type") == "text" for b in content if isinstance(b, dict)) and not any(
            b.get("type") == "tool_result" for b in content if isinstance(b, dict)
        )
    return False


def text_of(msg: dict) -> str:
    content = (msg.get("message") or {}).get("content")
    if isinstance(content, str):
        return content
    if isinstance(content, list):
        return "\n".join(b.get("text", "") for b in content if isinstance(b, dict))
    return ""


def skills_in_assistant(msg: dict) -> list[str]:
    out = []
    for b in (msg.get("message") or {}).get("content") or []:
        if isinstance(b, dict) and b.get("type") == "tool_use" and b.get("name") == "Skill":
            inp = b.get("input") or {}
            name = inp.get("skill") or inp.get("command") or inp.get("name")
            if name:
                out.append(str(name).lstrip("/"))
    return out


def parse_session(path: Path, pricing: dict) -> dict:
    turns, cur = [], None
    seen = set()
    first_ts = last_ts = None
    with open(path, encoding="utf-8") as f:
        for line in f:
            try:
                msg = json.loads(line)
            except json.JSONDecodeError:
                continue
            ts = msg.get("timestamp")
            if ts:
                first_ts = first_ts or ts
                last_ts = ts
            t = msg.get("type")
            if t == "user" and is_real_prompt(msg) and not msg.get("isSidechain"):
                txt = text_of(msg)
                shown = CMD_RE.sub(lambda mm: "/" + mm.group(1), txt)
                shown = re.sub(r"<[^>]+>", " ", shown)
                cur = {"prompt": " ".join(shown.split())[:80], "ts": ts,
                       "skills": set(CMD_RE.findall(txt)), "usage": defaultdict(int),
                       "cost": 0.0, "calls": 0, "models": set(), "unpriced": set()}
                turns.append(cur)
                continue
            if t != "assistant":
                continue
            m = msg.get("message") or {}
            # streaming writes one line per content block with the same usage -> dedupe
            key = (m.get("id"), msg.get("requestId"))
            if cur is None:
                cur = {"prompt": "(before first prompt)", "ts": ts, "skills": set(),
                       "usage": defaultdict(int), "cost": 0.0, "calls": 0,
                       "models": set(), "unpriced": set()}
                turns.append(cur)
            cur["skills"].update(skills_in_assistant(msg))
            if key in seen or not m.get("usage") or m.get("model") == "<synthetic>":
                continue
            seen.add(key)
            u = norm_usage(m["usage"])
            price = price_for_model(m.get("model"), pricing)
            c = cost_usd(u, price, pricing)
            for k in KEYS:
                cur["usage"][k] += u[k]
            cur["calls"] += 1
            cur["models"].add(m.get("model"))
            if c is None:
                cur["unpriced"].add(m.get("model"))
            else:
                cur["cost"] += c
    total = defaultdict(int)
    for tr in turns:
        for k in KEYS:
            total[k] += tr["usage"][k]
    return {"file": path, "id": path.stem, "project": path.parent.name,
            "start": first_ts, "end": last_ts, "turns": turns, "usage": total,
            "cost": sum(t["cost"] for t in turns),
            "unpriced": sorted({m for t in turns for m in t["unpriced"] if m})}


def total_tokens(u: dict) -> int:
    return sum(u[k] for k in KEYS)


def row(label, u, cost, fx):
    return [label, u["input"], u["output"], u["cache_write_5m"] + u["cache_write_1h"],
            u["cache_read"], total_tokens(u), fmt_usd(cost), fmt_thb(cost, fx)]


HEAD = ["item", "input", "output", "cache_write", "cache_read", "total", "USD", "THB"]


def md_table(rows, head=HEAD):
    lines = ["| " + " | ".join(head) + " |", "|" + "---|" * len(head)]
    for r in rows:
        lines.append("| " + " | ".join(f"{x:,}" if isinstance(x, int) else str(x) for x in r) + " |")
    return "\n".join(lines)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--path", help="projects dir or a single .jsonl file")
    ap.add_argument("--session", help="session id (or part of it) or .jsonl path")
    ap.add_argument("--all", action="store_true", help="all sessions instead of latest")
    ap.add_argument("--project", help="filter by project folder name substring")
    ap.add_argument("--since", help="YYYY-MM-DD, keep sessions modified on/after")
    ap.add_argument("--turns", action="store_true", help="show per-turn breakdown")
    ap.add_argument("--by-skill", action="store_true", help="aggregate turns by skill")
    ap.add_argument("--fx", type=float, help="USD->THB rate to also show baht")
    ap.add_argument("--pricing", help="custom pricing.json")
    ap.add_argument("--xlsx", help="write Excel report")
    ap.add_argument("--json", help="write raw JSON report")
    a = ap.parse_args()
    pricing = load_pricing(a.pricing)

    root = Path(a.path).expanduser() if a.path else projects_dir()
    if a.session and Path(a.session).expanduser().is_file():
        files = [Path(a.session).expanduser()]
    elif root.is_file():
        files = [root]
    else:
        if not root.exists():
            sys.exit(f"ไม่พบโฟลเดอร์ transcript: {root}")
        files = sorted(root.rglob("*.jsonl"), key=lambda p: p.stat().st_mtime)
        if a.project:
            files = [p for p in files if a.project.lower() in str(p.parent).lower()]
        if a.session:
            files = [p for p in files if a.session in p.stem]
        if a.since:
            cut = datetime.fromisoformat(a.since).timestamp()
            files = [p for p in files if p.stat().st_mtime >= cut]
        if not (a.all or a.since or a.project or a.session):
            files = files[-1:]
    if not files:
        sys.exit("ไม่พบ session ที่ตรงเงื่อนไข")

    sessions = [parse_session(p, pricing) for p in files]
    out = []

    grand = defaultdict(int)
    srows = []
    for s in sessions:
        for k in KEYS:
            grand[k] += s["usage"][k]
        srows.append(row(f"{s['id'][:8]} ({s['project'][-30:]})", s["usage"], s["cost"], a.fx))
    gcost = sum(s["cost"] for s in sessions)
    srows.append(row("**TOTAL**", grand, gcost, a.fx))
    out.append(f"## Sessions ({len(sessions)})\n\n" + md_table(srows))

    trows = []
    if a.turns:
        for s in sessions:
            for i, t in enumerate(s["turns"], 1):
                label = f"{s['id'][:8]} #{i} [{', '.join(sorted(t['skills'])) or '-'}] {t['prompt'][:40]}"
                trows.append(row(label, t["usage"], t["cost"], a.fx))
        out.append("## Turns\n\n" + md_table(trows))

    krows = []
    if a.by_skill:
        agg = defaultdict(lambda: {"usage": defaultdict(int), "cost": 0.0, "n": 0})
        for s in sessions:
            for t in s["turns"]:
                for sk in (t["skills"] or {"(no skill)"}):
                    g = agg[sk]
                    g["n"] += 1
                    g["cost"] += t["cost"]
                    for k in KEYS:
                        g["usage"][k] += t["usage"][k]
        for sk, g in sorted(agg.items(), key=lambda kv: -kv[1]["cost"]):
            r = row(sk, g["usage"], g["cost"], a.fx)
            r.insert(1, g["n"])
            r.append(fmt_usd(g["cost"] / g["n"]) if g["n"] else "")
            krows.append(r)
        head = HEAD[:1] + ["turns"] + HEAD[1:] + ["USD/turn"]
        out.append("## By skill (turns that invoked it)\n\n" + md_table(krows, head))

    unpriced = sorted({m for s in sessions for m in s["unpriced"]})
    if unpriced:
        out.append("⚠️ ไม่มีราคาในตาราง (cost ไม่นับรวม): " + ", ".join(unpriced))
    out.append("_cost = API list price equivalent. ผู้ใช้ Pro/Max จ่ายแบบเหมา ตัวเลขนี้ใช้เทียบขนาดการใช้งานเท่านั้น_")
    print("\n\n".join(out))

    if a.json:
        dump = [{**s, "file": str(s["file"]),
                 "turns": [{**t, "skills": sorted(t["skills"]), "models": sorted(filter(None, t["models"])),
                            "unpriced": sorted(filter(None, t["unpriced"]))}
                           for t in s["turns"]]} for s in sessions]
        Path(a.json).write_text(json.dumps(dump, ensure_ascii=False, indent=2, default=dict), encoding="utf-8")
    if a.xlsx:
        write_xlsx(a.xlsx, srows, trows, krows)


def write_xlsx(path, srows, trows, krows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    sheets = [("Sessions", HEAD, srows), ("Turns", HEAD, trows),
              ("BySkill", HEAD[:1] + ["turns"] + HEAD[1:] + ["USD/turn"], krows)]
    first = True
    for name, head, rows in sheets:
        if not rows:
            continue
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name
        ws.append(head)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78")
        for r in rows:
            ws.append([str(x).replace("**", "") if isinstance(x, str) else x for x in r])
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = max(12, min(60, max(len(str(c.value or "")) for c in col) + 2))
        for r in ws.iter_rows(min_row=2):
            for c in r:
                if isinstance(c.value, int):
                    c.number_format = "#,##0"
    wb.save(path)
    print(f"\nsaved: {path}")


if __name__ == "__main__":
    main()
